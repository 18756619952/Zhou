#!/usr/bin/env python3
"""将 PPTX 转为可展示的单页网页（HTML + 资源文件）。"""

from __future__ import annotations

import argparse
import html
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


CSS = """
:root {
  --bg: #0f172a;
  --panel: #111827;
  --text: #e5e7eb;
  --muted: #94a3b8;
  --accent: #38bdf8;
}
* { box-sizing: border-box; }
body {
  margin: 0;
  background: linear-gradient(135deg, #020617, #0f172a);
  color: var(--text);
  font-family: "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
}
.app {
  min-height: 100vh;
  display: flex;
  flex-direction: column;
}
.toolbar {
  display: flex;
  align-items: center;
  justify-content: space-between;
  padding: 14px 20px;
  background: rgba(17, 24, 39, 0.88);
  backdrop-filter: blur(10px);
  position: sticky;
  top: 0;
  z-index: 20;
  border-bottom: 1px solid rgba(148, 163, 184, 0.2);
}
.title { font-weight: 700; }
.controls { display: flex; gap: 10px; align-items: center; }
button {
  border: 1px solid rgba(148, 163, 184, 0.4);
  color: var(--text);
  background: rgba(30, 41, 59, 0.9);
  border-radius: 8px;
  padding: 8px 12px;
  cursor: pointer;
}
button:hover { border-color: var(--accent); }
.counter { color: var(--muted); min-width: 64px; text-align: center; }
.viewer {
  width: min(1100px, 92vw);
  margin: 18px auto 24px;
}
.slide {
  display: none;
  background: rgba(17, 24, 39, 0.95);
  border: 1px solid rgba(148, 163, 184, 0.2);
  border-radius: 14px;
  padding: 24px;
  min-height: 72vh;
  box-shadow: 0 30px 55px rgba(0, 0, 0, 0.35);
}
.slide.active { display: block; }
.slide-title {
  margin-top: 0;
  margin-bottom: 20px;
  color: #f8fafc;
}
.bullets { margin: 0; padding-left: 24px; line-height: 1.7; }
.text-block {
  white-space: pre-wrap;
  line-height: 1.8;
  color: #e2e8f0;
  margin-bottom: 14px;
}
.image-grid {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(240px, 1fr));
  gap: 12px;
  margin-top: 18px;
}
.image-grid img {
  width: 100%;
  border-radius: 10px;
  border: 1px solid rgba(148, 163, 184, 0.3);
  background: #020617;
}
.hint {
  color: var(--muted);
  font-size: 13px;
}
""".strip()


JS = """
const slides = Array.from(document.querySelectorAll('.slide'));
const counter = document.querySelector('#counter');
let index = 0;

function render() {
  slides.forEach((slide, i) => slide.classList.toggle('active', i === index));
  counter.textContent = `${index + 1} / ${slides.length}`;
}

function next() {
  index = (index + 1) % slides.length;
  render();
}

function prev() {
  index = (index - 1 + slides.length) % slides.length;
  render();
}

document.querySelector('#next').addEventListener('click', next);
document.querySelector('#prev').addEventListener('click', prev);
document.addEventListener('keydown', (event) => {
  if (event.key === 'ArrowRight' || event.key === 'PageDown') next();
  if (event.key === 'ArrowLeft' || event.key === 'PageUp') prev();
});

render();
""".strip()


def safe_stem(path: Path) -> str:
  return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in path.stem)


def text_lines(shape) -> list[str]:
  if not getattr(shape, "has_text_frame", False):
    return []
  lines: list[str] = []
  for paragraph in shape.text_frame.paragraphs:
    text = paragraph.text.strip()
    if text:
      lines.append(text)
  return lines


def save_picture(shape, output_images: Path, slide_i: int, image_i: int) -> str:
  image = shape.image
  ext = image.ext
  image_name = f"slide-{slide_i:02d}-image-{image_i:02d}.{ext}"
  image_path = output_images / image_name
  image_path.write_bytes(image.blob)
  return f"images/{image_name}"


def render_slide(slide_no: int, title: str, bullets: Iterable[str], texts: Iterable[str], images: Iterable[str]) -> str:
  title_html = f"<h2 class=\"slide-title\">{html.escape(title or f'Slide {slide_no}')}</h2>"

  bullet_items = "".join(f"<li>{html.escape(line)}</li>" for line in bullets)
  bullets_html = f"<ul class=\"bullets\">{bullet_items}</ul>" if bullet_items else ""

  text_html = "".join(f"<div class=\"text-block\">{html.escape(line)}</div>" for line in texts)
  image_html = "".join(f"<img src=\"{src}\" alt=\"slide image\" loading=\"lazy\">" for src in images)
  image_section = f"<div class=\"image-grid\">{image_html}</div>" if image_html else ""

  return f"""
<section class="slide">
  {title_html}
  {bullets_html}
  {text_html}
  {image_section}
</section>
""".strip()


def convert(pptx_path: Path, output_dir: Path) -> Path:
  prs = Presentation(pptx_path)

  output_dir.mkdir(parents=True, exist_ok=True)
  image_dir = output_dir / "images"
  image_dir.mkdir(parents=True, exist_ok=True)

  slides_html: list[str] = []

  for s_idx, slide in enumerate(prs.slides, start=1):
    title = ""
    bullets: list[str] = []
    texts: list[str] = []
    images: list[str] = []

    for shp_idx, shape in enumerate(slide.shapes, start=1):
      if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        images.append(save_picture(shape, image_dir, s_idx, shp_idx))
        continue

      lines = text_lines(shape)
      if not lines:
        continue

      if shape == slide.shapes.title and lines:
        title = lines[0]
        if len(lines) > 1:
          texts.extend(lines[1:])
      elif len(lines) > 1:
        bullets.extend(lines)
      else:
        texts.extend(lines)

    slides_html.append(render_slide(s_idx, title, bullets, texts, images))

  doc_title = html.escape(pptx_path.stem)
  output_html = f"""
<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{doc_title}</title>
  <style>{CSS}</style>
</head>
<body>
  <div class="app">
    <header class="toolbar">
      <div>
        <div class="title">{doc_title}</div>
        <div class="hint">键盘左右方向键 / PageUp / PageDown 可切换页面</div>
      </div>
      <div class="controls">
        <button id="prev" type="button">上一页</button>
        <div id="counter" class="counter"></div>
        <button id="next" type="button">下一页</button>
      </div>
    </header>
    <main class="viewer">
      {''.join(slides_html)}
    </main>
  </div>
  <script>{JS}</script>
</body>
</html>
""".strip()

  index_path = output_dir / "index.html"
  index_path.write_text(output_html, encoding="utf-8")
  return index_path


def parse_args() -> argparse.Namespace:
  parser = argparse.ArgumentParser(description="将 PPTX 转换为可展示网页")
  parser.add_argument("pptx", type=Path, help="输入 .pptx 文件路径")
  parser.add_argument(
    "-o",
    "--output",
    type=Path,
    default=Path("web_ppt"),
    help="输出目录（默认：web_ppt）",
  )
  return parser.parse_args()


def main() -> None:
  args = parse_args()
  pptx_path = args.pptx
  if not pptx_path.exists() or pptx_path.suffix.lower() != ".pptx":
    raise SystemExit("请提供存在的 .pptx 文件路径")

  output_dir = args.output / safe_stem(pptx_path)
  index_path = convert(pptx_path, output_dir)
  print(f"转换完成：{index_path}")
  print("可执行以下命令进行本地预览：")
  print(f"  python3 -m http.server 8000 -d {output_dir.parent}")


if __name__ == "__main__":
  main()
